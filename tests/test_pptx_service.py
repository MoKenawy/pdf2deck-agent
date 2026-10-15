from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

from pdf2deck.schemas import (
    DeckOutline,
    Fonts,
    ImageBox,
    Palette,
    SlidePlan,
    SlideSize,
    SourceImage,
    StyledSlide,
    StylePreferences,
    TypeScale,
)
from pdf2deck.services.pptx_service import PptxService

STYLE = StylePreferences(
    name="brand",
    slide_size=SlideSize(width_in=10, height_in=5.62),
    colors=Palette(main="#32AEB8", highlight="#F2A40D", text="#222222", background="#FAFAFA"),
    fonts=Fonts(heading="Georgia", body="Arial"),
    type_scale=TypeScale(title_pt=24, body_pt=12),
)


def _render(
    tmp_path: Path, style: StylePreferences = STYLE, image: SourceImage | None = None
) -> Presentation:
    plan = SlidePlan(
        slide_id="s1",
        title="First slide",
        purpose="Introduce the topic",
        key_points=["Point one", "Point two"],
        source_pages=[1],
        image_id=image.id if image else None,
        speaker_notes="Say this.",
    )
    outline = DeckOutline(title="Fake deck", subtitle="A subtitle", slides=[plan])
    styled = [
        StyledSlide(plan=plan, layout_name=plan.visual_type, style_name=style.name, image=image)
    ]

    artifact = PptxService().render(outline, styled, style, tmp_path / "out" / "deck.pptx")

    assert artifact.slide_count == 2
    assert artifact.title == "Fake deck"
    return Presentation(artifact.path)


def _runs(slide):
    return [
        run
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    ]


def test_slide_size_comes_from_the_preferences(tmp_path: Path):
    prs = _render(tmp_path)

    assert prs.slide_width == Inches(10)
    assert prs.slide_height == Inches(5.62)


def test_cover_slide_shows_title_and_subtitle(tmp_path: Path):
    cover = _render(tmp_path).slides[0]

    assert [run.text for run in _runs(cover)] == ["Fake deck", "A subtitle"]


def test_content_slide_applies_fonts_sizes_and_colors(tmp_path: Path):
    slide = _render(tmp_path).slides[1]
    title, *points = _runs(slide)

    assert title.text == "First slide"
    assert title.font.name == "Georgia"
    assert title.font.size == Pt(24)
    assert title.font.color.rgb == RGBColor.from_string("32AEB8")

    assert [point.text for point in points] == ["Point one", "Point two"]
    for point in points:
        assert point.font.name == "Arial"
        assert point.font.size == Pt(12)
        assert point.font.color.rgb == RGBColor.from_string("222222")


def test_background_and_rule_use_the_palette(tmp_path: Path):
    slide = _render(tmp_path).slides[1]
    rule = next(shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE)

    assert slide.background.fill.fore_color.rgb == RGBColor.from_string("FAFAFA")
    assert rule.fill.fore_color.rgb == RGBColor.from_string("F2A40D")


def _source_image(tmp_path: Path, width_px: int, height_px: int) -> SourceImage:
    path = tmp_path / "figure.png"
    Image.new("RGB", (width_px, height_px), "red").save(path)
    return SourceImage(
        id="img-p1-0",
        page=1,
        path=str(path),
        width_px=width_px,
        height_px=height_px,
        box=ImageBox(x=0, y=0, width=width_px, height=height_px),
        page_width_pt=612,
        page_height_pt=792,
    )


def _body(slide):
    return next(
        shape for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.startswith("Point one")
    )


def test_slide_image_sits_right_of_the_bullets_at_its_own_aspect_ratio(tmp_path: Path):
    prs = _render(tmp_path, image=_source_image(tmp_path, 400, 300))
    slide = prs.slides[1]
    [picture] = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    body = _body(slide)

    assert picture.left >= body.left + body.width
    assert picture.left + picture.width <= prs.slide_width
    assert picture.width / picture.height == pytest.approx(400 / 300, rel=0.01)
    assert body.width < _body(_render(tmp_path / "plain").slides[1]).width


def test_slide_without_an_image_has_no_picture(tmp_path: Path):
    slide = _render(tmp_path).slides[1]

    assert not any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)


def test_missing_image_file_is_reported(tmp_path: Path):
    image = _source_image(tmp_path, 400, 300)
    Path(image.path).unlink()

    with pytest.raises(ValueError, match="img-p1-0 is missing"):
        _render(tmp_path, image=image)


@pytest.mark.parametrize("size", [None, (400, 300), (200, 900)])
def test_everything_stays_inside_the_slide(tmp_path: Path, size: tuple[int, int] | None):
    prs = _render(tmp_path, image=_source_image(tmp_path, *size) if size else None)

    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width
            assert shape.top + shape.height <= prs.slide_height


def test_speaker_notes_follow_the_preference(tmp_path: Path):
    with_notes = _render(tmp_path).slides[1]
    without_notes = _render(tmp_path, STYLE.model_copy(update={"include_speaker_notes": False})).slides[1]

    assert with_notes.notes_slide.notes_text_frame.text == "Say this."
    assert not without_notes.has_notes_slide
