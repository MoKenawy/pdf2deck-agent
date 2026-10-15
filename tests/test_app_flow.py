from pathlib import Path

import gradio as gr
import pytest
from pptx import Presentation
from pptx.util import Inches

from pdf2deck.config import Settings
from pdf2deck.graph import DeckGraph
from pdf2deck.preferences import load_preferences
from pdf2deck.schemas import DeckOutline, SlidePlan, SourceDocument, SourcePage


class FakePdf:
    def extract_document(self, path: Path) -> SourceDocument:
        return SourceDocument(filename=path.name, pages=[SourcePage(number=1, text="Source text.")])


class FakeLlm:
    def create_outline(self, document, tone, density, max_slides) -> DeckOutline:
        slide = SlidePlan(
            slide_id="s1",
            title="First slide",
            purpose="Introduce the topic",
            key_points=["Point one", "Point two"],
            source_pages=[1],
        )
        return DeckOutline(title="Fake deck", learning_objectives=["Learn it"], slides=[slide])


@pytest.fixture
def app(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("OPENAI_API_KEY", "test-key")
    import pdf2deck.app as app_module

    # Isolated from the developer's .env: outputs go to tmp_path and no default file applies.
    settings = Settings(output_dir=tmp_path / "outputs", default_preferences_path=None)
    deck = DeckGraph(settings)
    deck.pdf = FakePdf()
    deck.llm = FakeLlm()
    monkeypatch.setattr(app_module, "graph", deck.build())
    monkeypatch.setattr(app_module, "settings", settings)
    return app_module


@pytest.fixture
def pdf_path(tmp_path: Path) -> str:
    path = tmp_path / "chapter.pdf"
    path.write_bytes(b"%PDF-1.4")
    return str(path)


def _generate(app, pdf_path: str):
    views = list(app.run_deck(pdf_path, "technical-minimal", "educational", "balanced", "Aptos", None))
    return views[-1]


def _enabled(buttons) -> list[bool]:
    return [button.interactive for button in buttons]


def test_generate_shows_outline_and_waits_for_outline_review(app, pdf_path):
    run_id, status, outline, log, output, *buttons = _generate(app, pdf_path)

    assert run_id
    assert "Outline ready" in status
    assert "## Fake deck" in outline
    assert "### 1. First slide" in outline
    assert "- Point one" in outline
    assert "build_outline: Outline generated." in log
    assert output is None
    assert _enabled(buttons) == [True, True, False, False]


def test_generate_reports_the_running_step_before_it_finishes(app, pdf_path):
    views = list(app.run_deck(pdf_path, "technical-minimal", "educational", "balanced", "Aptos", None))

    during_outline = next(view for view in views if "[started] build_outline" in view[3])
    assert "Generating the outline" in during_outline[1]
    assert "extract_pdf: Extracted 1 pages" in during_outline[3]
    assert "Outline generated" not in during_outline[3]
    assert "[waiting] outline_review" in views[-1][3]


def test_model_failure_is_reported_instead_of_leaving_the_run_spinning(app, pdf_path, monkeypatch):
    def fail(*args, **kwargs):
        raise RuntimeError("model unavailable")

    monkeypatch.setattr(FakeLlm, "create_outline", fail)

    _, status, _, log, output, *buttons = _generate(app, pdf_path)
    assert status == "**Failed:** RuntimeError: model unavailable"
    assert "extract_pdf: Extracted 1 pages" in log
    assert output is None
    assert _enabled(buttons) == [False, False, False, False]


def test_approving_outline_then_deck_returns_the_pptx(app, pdf_path):
    run_id = _generate(app, pdf_path)[0]

    status, outline, _, output, *buttons = list(app.resume_outline(run_id, True))[-1]
    assert "Deck rendered (2 slides)" in status
    assert "## Fake deck" in outline
    assert Path(output).exists()
    assert _enabled(buttons) == [False, False, True, True]

    status, _, log, output, *buttons = list(app.resume_deck(run_id, True))[-1]
    assert "Deck approved" in status
    assert "approve_deck: Deck approved." in log
    assert Path(output).exists()
    assert _enabled(buttons) == [False, False, False, False]


def test_deck_review_cannot_resume_a_pending_outline_review(app, pdf_path):
    run_id = _generate(app, pdf_path)[0]

    with pytest.raises(gr.Error, match="No deck is waiting"):
        list(app.resume_deck(run_id, True))


def _render_with(app, pdf_path: str, preferences_file: str | None) -> Presentation:
    views = list(
        app.run_deck(pdf_path, "technical-minimal", "educational", "balanced", "Aptos", preferences_file)
    )
    output = list(app.resume_outline(views[-1][0], True))[-1][3]
    return Presentation(output)


def _fonts(prs: Presentation) -> set[str]:
    return {
        run.font.name
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    }


def test_form_sets_the_style_when_no_preferences_file_is_given(app, pdf_path):
    assert _fonts(_render_with(app, pdf_path, None)) == {"Aptos"}


def test_uploaded_preferences_file_wins_over_the_form(app, pdf_path, tmp_path):
    preferences = tmp_path / "preferences.json"
    preferences.write_text(
        '{"fonts": {"heading": "Georgia", "body": "Arial"}, "slide_size": {"width_in": 10, "height_in": 5.62}}',
        encoding="utf-8",
    )

    prs = _render_with(app, pdf_path, str(preferences))

    assert _fonts(prs) == {"Georgia", "Arial"}
    assert prs.slide_width == Inches(10)


def test_configured_default_preferences_file_is_used_when_none_is_uploaded(
    app, pdf_path, tmp_path, monkeypatch
):
    default = tmp_path / "default.json"
    default.write_text('{"fonts": {"heading": "Verdana", "body": "Verdana"}}', encoding="utf-8")
    monkeypatch.setattr(app.settings, "default_preferences_path", default)

    assert _fonts(_render_with(app, pdf_path, None)) == {"Verdana"}
    assert all(not field["interactive"] for field in app.form_fields(None))


def test_invalid_preferences_file_stops_before_the_run_starts(app, pdf_path, tmp_path):
    preferences = tmp_path / "preferences.json"
    preferences.write_text('{"colors": {"main": "teal"}}', encoding="utf-8")

    with pytest.raises(gr.Error, match="colors.main"):
        list(app.run_deck(pdf_path, "technical-minimal", "educational", "balanced", "Aptos", str(preferences)))


def test_form_fields_lock_while_a_preferences_file_is_present(app):
    assert all(not field["interactive"] for field in app.form_fields("preferences.json"))
    assert all(field["interactive"] for field in app.form_fields(None))


def test_extract_button_writes_a_preferences_file_that_loads(app, tmp_path):
    template = tmp_path / "Brand Deck.pptx"
    Presentation().save(template)

    content, path = app.extract_preferences(str(template))

    assert Path(path).name == "brand-deck.preferences.json"
    assert Path(path).read_text(encoding="utf-8") == content
    assert load_preferences(Path(path)).name == "brand-deck"


def test_extract_button_reports_a_file_that_is_not_a_presentation(app, tmp_path):
    template = tmp_path / "broken.pptx"
    template.write_bytes(b"not a zip")

    with pytest.raises(gr.Error, match="Could not read preferences from broken.pptx"):
        app.extract_preferences(str(template))


def test_rejecting_outline_stops_the_run(app, pdf_path):
    run_id = _generate(app, pdf_path)[0]

    status, _, _, output, *buttons = list(app.resume_outline(run_id, False))[-1]
    assert "Outline was rejected" in status
    assert output is None
    assert _enabled(buttons) == [False, False, False, False]
