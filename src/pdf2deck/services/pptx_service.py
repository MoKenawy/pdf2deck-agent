from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from ..schemas import DeckArtifact, DeckOutline, StyledSlide, StylePreferences

# Index of the empty layout in python-pptx's bundled default template. Its other
# layouts position their placeholders for a 10 x 7.5 in slide, so every slide is
# drawn on the blank one from the preferences instead.
BLANK_LAYOUT = 6

# Slide geometry as fractions of the slide width and height.
MARGIN_X = 0.06
CONTENT_WIDTH = 1 - 2 * MARGIN_X
RULE_HEIGHT = 0.008
BULLET_INDENT = Inches(0.3)


class PptxService:
    def render(
        self,
        outline: DeckOutline,
        styled_slides: list[StyledSlide],
        style: StylePreferences,
        output_path: Path,
    ) -> DeckArtifact:
        prs = Presentation()
        prs.slide_width = Inches(style.slide_size.width_in)
        prs.slide_height = Inches(style.slide_size.height_in)
        blank = prs.slide_layouts[BLANK_LAYOUT]

        self._render_cover(prs, prs.slides.add_slide(blank), outline, style)
        for styled in styled_slides:
            self._render_slide(prs, prs.slides.add_slide(blank), styled, style)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(output_path)
        return DeckArtifact(
            path=str(output_path),
            slide_count=len(prs.slides),
            title=outline.title,
        )

    def _render_cover(
        self,
        prs: Presentation,
        slide,
        outline: DeckOutline,
        style: StylePreferences,
    ) -> None:
        self._fill_background(slide, style)

        title = self._text_frame(prs, slide, top=0.28, height=0.26)
        title.vertical_anchor = MSO_ANCHOR.BOTTOM
        paragraph = title.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        self._write(
            paragraph, outline.title, style.fonts.heading, style.type_scale.title_pt,
            style.colors.main, bold=True,
        )

        self._add_rule(prs, slide, style, left=0.4, top=0.57, width=0.2)

        subtitle = self._text_frame(prs, slide, top=0.61, height=0.2)
        paragraph = subtitle.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        self._write(
            paragraph, outline.subtitle or "Generated from source chapter",
            style.fonts.body, style.type_scale.body_pt, style.colors.text,
        )

    def _render_slide(
        self,
        prs: Presentation,
        slide,
        styled: StyledSlide,
        style: StylePreferences,
    ) -> None:
        plan = styled.plan
        self._fill_background(slide, style)

        title = self._text_frame(prs, slide, top=0.06, height=0.16)
        title.vertical_anchor = MSO_ANCHOR.BOTTOM
        self._write(
            title.paragraphs[0], plan.title, style.fonts.heading, style.type_scale.title_pt,
            style.colors.main, bold=True,
        )

        self._add_rule(prs, slide, style, left=MARGIN_X, top=0.235, width=CONTENT_WIDTH)

        body = self._text_frame(prs, slide, top=0.27, height=0.65)
        for index, point in enumerate(plan.key_points):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.space_after = Pt(style.type_scale.body_pt * 0.4)
            self._add_bullet(paragraph)
            self._write(
                paragraph, point, style.fonts.body, style.type_scale.body_pt, style.colors.text
            )

        if style.include_speaker_notes and plan.speaker_notes:
            slide.notes_slide.notes_text_frame.text = plan.speaker_notes

    @staticmethod
    def _rgb(color: str) -> RGBColor:
        return RGBColor.from_string(color.lstrip("#").upper())

    def _fill_background(self, slide, style: StylePreferences) -> None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self._rgb(style.colors.background)

    def _text_frame(self, prs: Presentation, slide, top: float, height: float):
        box = slide.shapes.add_textbox(
            int(prs.slide_width * MARGIN_X),
            int(prs.slide_height * top),
            int(prs.slide_width * CONTENT_WIDTH),
            int(prs.slide_height * height),
        )
        frame = box.text_frame
        frame.word_wrap = True
        # No side insets, so text lines up with the rule.
        frame.margin_left = frame.margin_right = 0
        return frame

    def _add_rule(
        self, prs: Presentation, slide, style: StylePreferences,
        left: float, top: float, width: float,
    ) -> None:
        rule = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            int(prs.slide_width * left),
            int(prs.slide_height * top),
            int(prs.slide_width * width),
            int(prs.slide_height * RULE_HEIGHT),
        )
        rule.fill.solid()
        rule.fill.fore_color.rgb = self._rgb(style.colors.highlight)
        rule.line.fill.background()
        rule.shadow.inherit = False

    @staticmethod
    def _add_bullet(paragraph) -> None:
        # python-pptx has no bullet API, and text boxes do not bullet by default.
        # Must run after spacing is set: buChar follows spcAft in the pPr schema order.
        properties = paragraph._p.get_or_add_pPr()
        properties.set("marL", str(BULLET_INDENT))
        properties.set("indent", str(-BULLET_INDENT))
        properties.append(properties.makeelement(qn("a:buChar"), {"char": "•"}))

    def _write(
        self, paragraph, text: str, font: str, size_pt: float, color: str, bold: bool = False
    ) -> None:
        run = paragraph.add_run()
        run.text = text
        run.font.name = font
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.color.rgb = self._rgb(color)
