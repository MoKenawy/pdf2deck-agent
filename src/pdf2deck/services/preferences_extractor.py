import re
from collections import Counter
from pathlib import Path
from statistics import median
from typing import Literal, NamedTuple

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu

from ..schemas import Fonts, Palette, SlideSize, StylePreferences, TypeScale

ACCENTS = ("accent1", "accent2", "accent3", "accent4", "accent5", "accent6")
SYSTEM_COLORS = {"windowText": "000000", "window": "FFFFFF"}
# How a master maps the names shapes use onto theme slots when it has no clrMap.
DEFAULT_COLOR_MAP = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}
TITLE_PLACEHOLDERS = ("title", "ctrTitle")

# A title is the most common text size at least this many times the body size.
TITLE_RATIO = 1.5
# Median characters per slide below which a deck counts as sparse, then balanced.
SPARSE_BELOW = 250
BALANCED_BELOW = 600
# Channel spread (max - min, out of 255) from which a colour counts as a brand colour
# rather than a grey, slate or pale tint.
VIVID_CHROMA = 64

# Colours a slide sets directly on its shapes. These paths skip <p:style>, whose
# lnRef/fillRef name accent1 on every autoshape even when its own fill overrides it.
SHAPE_COLORS = (
    f".//{qn('p:spPr')}/{qn('a:solidFill')}/*",
    f".//{qn('p:spPr')}/{qn('a:ln')}/{qn('a:solidFill')}/*",
    f".//{qn('a:tcPr')}/{qn('a:solidFill')}/*",
)


class Run(NamedTuple):
    """One piece of slide text: whether it is set as a title, its colour and typeface."""

    title: bool
    color: str | None
    typeface: str


class PreferencesExtractor:
    """Reads a starter StylePreferences out of an existing presentation.

    Colours and fonts come from what the slides use, by the role the renderer gives
    them; the theme is the fallback for a deck that sets nothing directly.
    """

    def extract(self, path: Path) -> StylePreferences:
        try:
            return self._extract(path)
        except Exception as exc:
            # python-pptx, zipfile and lxml each raise their own errors for a broken file.
            raise ValueError(f"Could not read preferences from {path.name}: {exc}") from exc

    def _extract(self, path: Path) -> StylePreferences:
        prs = Presentation(str(path))
        master = prs.slide_masters[0]
        theme = parse_xml(master.part.part_related_by(RT.THEME).blob)
        slides = [slide._element for slide in prs.slides]

        colors = self._named_colors(theme, master._element)
        theme_fonts = self._fonts(theme)
        type_scale = self._type_scale(slides)
        runs = self._runs(slides, colors, theme_fonts, type_scale.body_pt * TITLE_RATIO)
        backgrounds = [self._background(slide, colors) for slide in prs.slides]

        return StylePreferences(
            name=self._name(path),
            density=self._density(slides),
            slide_size=self._slide_size(prs),
            colors=self._palette(colors, slides, runs, backgrounds),
            fonts=self._run_fonts(runs, theme_fonts),
            type_scale=type_scale,
        )

    @staticmethod
    def _name(path: Path) -> str:
        slug = re.sub(r"[^a-z0-9]+", "-", path.stem.lower()).strip("-")
        return slug[:100] or "extracted"

    @staticmethod
    def _slide_size(prs: Presentation) -> SlideSize:
        if prs.slide_width is None or prs.slide_height is None:
            return SlideSize()
        return SlideSize(
            width_in=round(Emu(prs.slide_width).inches, 2),
            height_in=round(Emu(prs.slide_height).inches, 2),
        )

    @staticmethod
    def _hex(value, colors: dict[str, str]) -> str | None:
        """RRGGBB for a colour element, looking scheme colours up in `colors`."""
        hex_value: str | None
        if value.tag == qn("a:srgbClr"):
            hex_value = value.get("val")
        elif value.tag == qn("a:schemeClr"):
            hex_value = colors.get(value.get("val", ""))
        elif value.tag == qn("a:sysClr"):
            hex_value = value.get("lastClr") or SYSTEM_COLORS.get(value.get("val", ""))
        else:
            hex_value = None
        if hex_value and re.fullmatch(r"[0-9A-Fa-f]{6}", hex_value):
            return hex_value.upper()
        return None

    def _scheme_colors(self, theme) -> dict[str, str]:
        """Map theme colour slots (dk1, lt1, accent1, ...) to RRGGBB."""
        colors: dict[str, str] = {}
        scheme = theme.find(f".//{qn('a:clrScheme')}")
        for slot in scheme if scheme is not None else ():
            hex_value = self._hex(slot[0], {})
            if hex_value:
                colors[slot.tag.rpartition("}")[2]] = hex_value
        return colors

    def _named_colors(self, theme, master) -> dict[str, str]:
        """Theme slots plus the names shapes use for them (tx1, bg1), mapped by the master."""
        scheme = self._scheme_colors(theme)
        color_map = master.find(qn("p:clrMap"))
        aliases = DEFAULT_COLOR_MAP | (dict(color_map.attrib) if color_map is not None else {})
        return scheme | {name: scheme[slot] for name, slot in aliases.items() if slot in scheme}

    def _background(self, slide, colors: dict[str, str]) -> str | None:
        """The slide's solid background, inherited from its layout and master if unset."""
        layout = slide.slide_layout
        for element in (slide._element, layout._element, layout.slide_master._element):
            background = element.find(f"{qn('p:cSld')}/{qn('p:bg')}")
            if background is None:
                continue
            value = background.find(f"{qn('p:bgPr')}/{qn('a:solidFill')}/*")
            if value is None:
                value = background.find(f"{qn('p:bgRef')}/*")
            # A gradient or picture background has no single colour.
            return self._hex(value, colors) if value is not None else None
        return colors.get("bg1")

    def _runs(
        self, slides: list, colors: dict[str, str], theme_fonts: Fonts, title_from_pt: float
    ) -> list[Run]:
        runs: list[Run] = []
        for slide in slides:
            for run in slide.iter(qn("a:r")):
                if not (run.findtext(qn("a:t")) or "").strip():
                    continue
                shape = next(run.iterancestors(qn("p:sp")), None)
                placeholder = (
                    shape.find(f"{qn('p:nvSpPr')}/{qn('p:nvPr')}/{qn('p:ph')}")
                    if shape is not None else None
                )
                in_title = placeholder is not None and placeholder.get("type") in TITLE_PLACEHOLDERS

                properties = run.find(qn("a:rPr"))
                size = properties.get("sz") if properties is not None else None
                value = (
                    properties.find(f"{qn('a:solidFill')}/*") if properties is not None else None
                )
                latin = properties.find(qn("a:latin")) if properties is not None else None
                typeface = latin.get("typeface", "") if latin is not None else ""
                if typeface == "+mj-lt" or (not typeface and in_title):
                    typeface = theme_fonts.heading
                elif typeface == "+mn-lt" or not typeface:
                    typeface = theme_fonts.body

                runs.append(Run(
                    # Sizes are stored in hundredths of a point.
                    title=in_title or (size is not None and int(size) / 100 >= title_from_pt),
                    # Text without its own colour inherits the theme's text colour.
                    color=self._hex(value, colors) if value is not None else colors.get("tx1"),
                    typeface=typeface,
                ))
        return runs

    @staticmethod
    def _vivid(color: str) -> bool:
        channels = bytes.fromhex(color)
        return max(channels) - min(channels) >= VIVID_CHROMA

    def _palette(
        self,
        colors: dict[str, str],
        slides: list,
        runs: list[Run],
        backgrounds: list[str | None],
    ) -> Palette:
        used: Counter[str] = Counter(run.color for run in runs if run.color)
        for slide in slides:
            for path in SHAPE_COLORS:
                for value in slide.iterfind(path):
                    if hex_value := self._hex(value, colors):
                        used[hex_value] += 1

        # Main is the brand colour the titles use, else the one the slides use most;
        # highlight is the next. Theme accents only stand in when the slides have none.
        titles = Counter(run.color for run in runs if run.title and run.color)
        title_vivid = [color for color, _ in titles.most_common() if self._vivid(color)]
        vivid = [color for color, _ in used.most_common() if self._vivid(color)]
        ranked = list(dict.fromkeys(title_vivid[:1] + vivid))
        if not ranked:
            # Distinct accents by how often the slides use them; ties keep theme order.
            accents = list(dict.fromkeys(colors[slot] for slot in ACCENTS if slot in colors))
            ranked = sorted(accents, key=lambda color: -used[color])

        defaults = Palette()
        main = f"#{ranked[0]}" if ranked else defaults.main
        highlight = f"#{ranked[1]}" if len(ranked) > 1 else main

        body = Counter(
            run.color for run in runs
            if not run.title and run.color and not self._vivid(run.color)
        )
        text = body.most_common(1)[0][0] if body else colors.get("tx1")
        filled = Counter(color for color in backgrounds if color)
        background = filled.most_common(1)[0][0] if filled else colors.get("bg1")
        return Palette(
            main=main,
            highlight=highlight,
            text=f"#{text}" if text else defaults.text,
            background=f"#{background}" if background else defaults.background,
        )

    @staticmethod
    def _fonts(theme) -> Fonts:
        defaults = Fonts()
        scheme = theme.find(f".//{qn('a:fontScheme')}")

        def typeface(kind: str, default: str) -> str:
            latin = scheme.find(f"{qn(kind)}/{qn('a:latin')}") if scheme is not None else None
            return (latin.get("typeface") if latin is not None else "") or default

        return Fonts(
            heading=typeface("a:majorFont", defaults.heading),
            body=typeface("a:minorFont", defaults.body),
        )

    @staticmethod
    def _run_fonts(runs: list[Run], theme_fonts: Fonts) -> Fonts:
        """The typefaces the titles and the body text use most, else the theme's."""

        def most_common(title: bool, default: str) -> str:
            faces = Counter(run.typeface for run in runs if run.title == title and run.typeface)
            return faces.most_common(1)[0][0][:100] if faces else default

        return Fonts(
            heading=most_common(True, theme_fonts.heading),
            body=most_common(False, theme_fonts.body),
        )

    @staticmethod
    def _type_scale(slides: list) -> TypeScale:
        sizes: Counter[float] = Counter()
        for slide in slides:
            for run in slide.iter(qn("a:rPr")):
                if run.get("sz"):
                    # Sizes are stored in hundredths of a point.
                    sizes[min(max(int(run.get("sz")) / 100, 6), 200)] += 1
        if not sizes:
            return TypeScale()

        defaults = TypeScale()
        body = sizes.most_common(1)[0][0]
        titles = [size for size, _ in sizes.most_common() if size >= body * TITLE_RATIO]
        title = titles[0] if titles else min(body * defaults.title_pt / defaults.body_pt, 200)
        return TypeScale(title_pt=round(title, 1), body_pt=body)

    @staticmethod
    def _density(slides: list) -> Literal["sparse", "balanced", "dense"]:
        if not slides:
            return StylePreferences().density
        characters = median(
            sum(len(text.text or "") for text in slide.iter(qn("a:t"))) for slide in slides
        )
        if characters < SPARSE_BELOW:
            return "sparse"
        return "balanced" if characters < BALANCED_BELOW else "dense"
